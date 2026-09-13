"""Synthetic XLSX intake: safe literal values, immutable evidence and reuse."""
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import csv
import json
import pytest
from openpyxl import Workbook
from utilityos.db import Store
from utilityos.operations import backup, check, restore, instance_lock
from utilityos.service import Ledger
from utilityos.spreadsheet import workbook
from utilityos.usage import UsageImport
from test_usage import RAW, MAPPING, usage, approve

REGION = {'sheet':'Water export','header_row':3,'first_column':2,'last_column':7,'end_row':6}


def xlsx(raw=RAW, alter=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Water export'
    sheet['B1'] = 'Fictional local meter export — no provider compatibility claim'
    for number, row in enumerate(csv.reader(BytesIO(raw).read().decode().splitlines()),3):
        for column,value in enumerate(row,2):
            sheet.cell(number,column,value)
    if alter:
        alter(book)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def rewrite(raw,name,transform):
    output=BytesIO()
    with ZipFile(BytesIO(raw)) as old, ZipFile(output,'w',ZIP_DEFLATED) as new:
        for entry in old.infolist():
            value=old.read(entry)
            new.writestr(entry.filename,transform(value) if entry.filename==name else value)
    return output.getvalue()


def ready_xlsx(usage,raw=None,region=None,mapping=None,reuse=True):
    identifier=usage.import_file('fictional.xlsx',raw or xlsx())['import_id']
    usage.preview(identifier,{'revision':0,'mapping':mapping or MAPPING,'source_region':region or REGION,'reuse_layout':reuse})
    return identifier


def test_workbook_retention_region_cell_provenance_and_financial_separation(usage,ledger):
    raw=xlsx()
    before=ledger.overview()
    identifier=usage.import_file('fictional.xlsx',raw)['import_id']
    info=usage.detail(identifier)
    assert info['source_region'] is None and info['headers']==[]
    assert info['sheets'][0]['name']=='Water export'
    assert info['sheets'][0]['candidate_headers'][0]['header_row']==3
    info=usage.detail(identifier,REGION)
    assert info['row_count']==3 and info['source_rows'][0]['value']=='12.50'
    usage.preview(identifier,{'revision':0,'mapping':MAPPING,'source_region':REGION,'reuse_layout':True})
    info=usage.detail(identifier)
    provenance=info['preview']['provenance']
    assert provenance['rows'][0]['value_column']['cell']=='E4'
    assert provenance['rows'][0]['value_column']['raw_value']=='12.50'
    assert provenance['rows'][0]['value_column']['normalized_value']=='12.5'
    assert provenance['rows'][0]['start_column']['normalized_value']==1785556800
    assert provenance['unit']=='US_gal' and provenance['region']==REGION
    assert info['preview']['mapping_version']=={'import_id':identifier,'revision':1}
    assert approve(usage,identifier)['reading_count']==3
    assert ledger.overview()==before
    assert (usage.store.sources/(sha256(raw).hexdigest()+'.xlsx')).read_bytes()==raw
    assert usage.import_file('renamed.xlsx',raw)=={'import_id':identifier,'duplicate_source':True}
    assert check(usage.store)['sources']=='ok'


def test_reuse_layout_after_approval_reuses_building_meter_identity_across_restart(usage):
    with usage.store.connect() as db:
        building=db.execute("INSERT INTO buildings(name) VALUES ('Fictional North Building')").lastrowid
        db.execute("UPDATE meters SET building_id=? WHERE code='SYN-W'",(building,))
    first=ready_xlsx(usage)
    future=xlsx(RAW.replace(b'2026-08-01',b'2026-09-01'))
    second=usage.import_file('later.xlsx',future)['import_id']
    assert usage.detail(second)['suggested_mapping'] is None # Pending mapping is not authority.
    approve(usage,first)
    restarted=UsageImport(Ledger(usage.store))
    info=restarted.detail(second)
    assert info['suggested_mapping']['mapping']==MAPPING
    assert info['suggested_mapping']['meter_reused'] is True
    assert info['source_region']==REGION
    usage.preview(second,{'revision':0,'mapping':info['suggested_mapping']['mapping'],'source_region':info['source_region'],'reuse_layout':True})
    approve(usage,second)
    with usage.store.connect() as db:
        assert db.execute("SELECT building_id FROM meters WHERE code='SYN-W'").fetchone()[0]==building
        assert db.execute('SELECT COUNT(DISTINCT meter_id) FROM usage_readings').fetchone()[0]==1


def test_same_layout_new_meter_requires_confirmation_then_relationship_reused(usage):
    first=ready_xlsx(usage)
    approve(usage,first)
    new_raw=RAW.replace(b'SYN-WATER-01',b'SYN-WATER-02')
    second=usage.import_file('second-building.xlsx',xlsx(new_raw))['import_id']
    suggestion=usage.detail(second)['suggested_mapping']
    assert suggestion['mapping']['meter_code']=='' and suggestion['meter_reused'] is False
    with usage.store.connect() as db:
        db.execute("INSERT INTO meters(code,commodity,unit) VALUES ('SYN-W2','water','gal')")
    mapping={**suggestion['mapping'],'meter_code':'SYN-W2'}
    usage.preview(second,{'revision':0,'mapping':mapping,'source_region':REGION,'reuse_layout':True})
    approve(usage,second)
    third=usage.import_file('second-building-next.xlsx',xlsx(new_raw.replace(b'2026-08-01',b'2026-09-01')))['import_id']
    assert usage.detail(third)['suggested_mapping']['mapping']==mapping
    assert usage.listing()['active_reading_count']==6


def test_reuse_extends_confirmed_tail_but_rejects_layout_or_unit_conflicts(usage):
    first=ready_xlsx(usage)
    approve(usage,first)
    raw=RAW.replace(b'2026-08-01',b'2026-09-01')
    rows=raw.splitlines()
    longer=b'\n'.join(rows+[rows[-1].replace(b'T02:',b'T03:').replace(b'T03:00:00-04:00,11',b'T04:00:00-04:00,11')])+b'\n'
    # An exact repeat row remains visible and is explicitly counted as a duplicate.
    longer=raw+rows[-1]+b'\n'
    identifier=usage.import_file('longer.xlsx',xlsx(longer))['import_id']
    assert usage.detail(identifier)['source_region']['end_row']==7
    changed=usage.import_file('changed-unit.xlsx',xlsx(raw.replace(b'US_gal',b'L')))['import_id']
    assert usage.detail(changed)['suggested_mapping'] is None
    changed=usage.import_file('changed-layout.xlsx',xlsx(raw.replace(b'source_meter,',b'new_meter_column,')))['import_id']
    assert usage.detail(changed)['suggested_mapping'] is None


def test_withdrawn_approved_layout_is_not_reused(usage):
    first=ready_xlsx(usage)
    approve(usage,first)
    usage.close(first,{'revision':1,'acknowledge':True,'reason':'Fictional incorrect meter relationship'},withdraw=True)
    identifier=usage.import_file('next.xlsx',xlsx(RAW.replace(b'2026-08-01',b'2026-09-01')))['import_id']
    assert usage.detail(identifier)['suggested_mapping'] is None


@pytest.mark.parametrize('alter,code',[
    (lambda b:setattr(b['Water export']['E4'],'value','=1+1'),'FORMULA'),
    (lambda b:setattr(b['Water export']['E4'],'value','  @SUM(A1)'),'FORMULA'),
    (lambda b:setattr(b['Water export']['E4'],'hyperlink','https://example.invalid'),'EXTERNAL'),
    (lambda b:setattr(b.create_sheet('Hidden'),'sheet_state','hidden'),'HIDDEN_SHEETS'),
    (lambda b:setattr(b['Water export']['E4'],'value',True),'BOOLEAN'),
    (lambda b:setattr(b['Water export']['C4'],'value',datetime(2026,8,1)) or setattr(b['Water export']['C4'],'number_format','yyyy-mm-dd'),'DATE_OR_TIME_AMBIGUOUS'),
    (lambda b:setattr(b['Water export']['C4'],'value',datetime(2026,8,1,0,0,0,500000)),'DATE_OR_TIME_AMBIGUOUS'),
    (lambda b:b['Water export'].cell(6001,1,'too far'),'RESOURCE_LIMIT'),
    (lambda b:b['Water export'].cell(1,51,'too wide'),'RESOURCE_LIMIT'),
])
def test_unsafe_or_ambiguous_workbooks_fail_before_retention(usage,alter,code):
    with pytest.raises(ValueError,match=code):
        usage.import_file('unsafe.xlsx',xlsx(alter=alter))
    assert usage.listing()['items']==[]


@pytest.mark.parametrize('alter',[
    lambda b:b['Water export'].merge_cells('B3:C3'),
    lambda b:setattr(b['Water export'].row_dimensions[4],'hidden',True),
    lambda b:setattr(b['Water export'].column_dimensions['E'],'hidden',True),
])
def test_selected_hidden_or_merged_cells_require_new_region(usage,alter):
    identifier=usage.import_file('region.xlsx',xlsx(alter=alter))['import_id']
    with pytest.raises(ValueError,match='REGION_CONTAINS_MERGED_OR_HIDDEN'):
        usage.preview(identifier,{'revision':0,'mapping':MAPPING,'source_region':REGION})
    assert usage.detail(identifier)['preview'] is None
    assert check(usage.store)['sources']=='ok'


def test_formula_cache_macro_external_refresh_and_dtd_rejected(usage):
    raw=xlsx()
    examples=[
        rewrite(raw,'xl/worksheets/sheet1.xml',lambda v:v.replace(b'<c r="E4" t="inlineStr"><is><t>12.50</t></is></c>',b'<c r="E4"><f>1+1</f><v>2</v></c>')),
        rewrite(raw,'[Content_Types].xml',lambda v:v.replace(b'spreadsheetml.sheet.main+xml',b'ms-excel.sheet.macroEnabled.main+xml')),
        rewrite(raw,'xl/workbook.xml',lambda v:b'<!DOCTYPE x [<!ENTITY x "unsafe">]>'+v),
    ]
    for raw in examples:
        with pytest.raises(ValueError,match='FORMULA|ACTIVE_OR_UNSUPPORTED|INVALID_OR_UNSUPPORTED'):
            usage.import_file('unsafe.xlsx',raw)
    raw=xlsx()
    output=BytesIO()
    with ZipFile(BytesIO(raw)) as old,ZipFile(output,'w',ZIP_DEFLATED) as new:
        for n in old.namelist():new.writestr(n,old.read(n))
        new.writestr('xl/connections.xml',b'<connections/>')
    with pytest.raises(ValueError,match='ACTIVE_OR_UNSUPPORTED'):
        usage.import_file('refresh.xlsx',output.getvalue())
    assert usage.listing()['items']==[]


def test_archive_bomb_duplicate_path_and_legacy_encryption_rejected(usage):
    for name,content,code in [('xl/huge.xml',b' '*2000000,'RESOURCE_LIMIT'),('../escape.xml',b'<x/>','UNSAFE')]:
        output=BytesIO()
        with ZipFile(output,'w',ZIP_DEFLATED) as new:new.writestr(name,content)
        with pytest.raises(ValueError,match=code):usage.import_file('unsafe.xlsx',output.getvalue())
    with pytest.raises(ValueError,match='ENCRYPTED_OR_LEGACY'):
        usage.import_file('encrypted.xlsx',b'\xd0\xcf\x11\xe0'+b'fake')
    with pytest.raises(ValueError,match='REQUIRES_CSV_OR_VALUES_ONLY_XLSX'):
        usage.import_file('old.xls',b'legacy')
    with pytest.raises(ValueError,match='INVALID_OR_UNSUPPORTED'):
        usage.import_file('partial.xlsx',xlsx()[:100])


def test_native_datetime_requires_timezone_and_preserves_xml_serial(usage):
    def native(b):
        for r in range(4,7):
            for c in ('C','D'):
                cell=b['Water export'][f'{c}{r}']
                cell.value=datetime.fromisoformat(cell.value).replace(tzinfo=None)
                cell.number_format='yyyy-mm-dd hh:mm:ss'
    raw=xlsx(alter=native)
    identifier=usage.import_file('native-datetime.xlsx',raw)['import_id']
    with pytest.raises(ValueError,match='NAIVE_TIMESTAMP_NEEDS_TIMEZONE'):
        usage.preview(identifier,{'revision':0,'mapping':MAPPING,'source_region':REGION})
    usage.preview(identifier,{'revision':0,'mapping':{**MAPPING,'timezone':'America/New_York'},'source_region':REGION})
    data=usage.detail(identifier)['preview']
    assert data['readings'][0]['start_utc']==1785556800
    assert data['provenance']['rows'][0]['start_column']['raw_value']=='46235'
    approve(usage,identifier)


def test_numeric_literals_use_exact_xml_decimal_instead_of_binary_rounding(usage):
    raw=xlsx(alter=lambda b:setattr(b['Water export']['E4'],'value',12.5))
    raw=rewrite(raw,'xl/worksheets/sheet1.xml',lambda v:v.replace(b'<v>12.5</v>',b'<v>1234567890123456.123456789</v>'))
    identifier=ready_xlsx(usage,raw)
    data=usage.detail(identifier)['preview']
    assert data['readings'][0]['quantity']=='1234567890123456.123456789'
    assert data['provenance']['rows'][0]['value_column']['raw_value']=='1234567890123456.123456789'


def test_xlsx_backup_restore_source_and_reattempt_preserve_history(usage,tmp_path):
    raw=xlsx()
    first=ready_xlsx(usage,raw)
    approve(usage,first)
    usage.close(first,{'revision':1,'acknowledge':True,'reason':'Fictional mapping corrected'},withdraw=True)
    next_id=usage.reattempt(first,{'revision':1,'acknowledge':True,'reason':'Review same workbook again'})['import_id']
    usage.preview(next_id,{'revision':0,'mapping':MAPPING,'source_region':REGION,'reuse_layout':True})
    approve(usage,next_id)
    saved=backup(usage.store)
    recovered=Store(tmp_path/'recovered','demo')
    with instance_lock(recovered.directory):restore(recovered,saved,'demo')
    restored=UsageImport(Ledger(recovered))
    assert restored.detail(first)['state']=='withdrawn'
    assert restored.detail(next_id)['source_history']['reattempt_of']==first
    assert restored.listing()['active_reading_count']==3
    assert (recovered.sources/(sha256(raw).hexdigest()+'.xlsx')).read_bytes()==raw
    assert check(recovered)['audit']=='ok'


def test_spreadsheet_api_inspection_approval_and_original_download(authenticated):
    client=authenticated
    with client.app.state.store.connect() as db:db.execute("INSERT INTO meters(code,commodity,unit) VALUES ('SYN-W','water','gal')")
    raw=xlsx()
    response=client.post('/api/usage/import',content=raw,headers={'content-type':'application/octet-stream','x-filename':'fictional.xlsx','x-synthetic-data':'true'})
    assert response.status_code==200,response.text
    identifier=response.json()['import_id']
    inspected=client.post(f'/api/usage/{identifier}/inspect',json={'source_region':REGION})
    assert inspected.status_code==200 and inspected.json()['row_count']==3
    assert client.post(f'/api/usage/{identifier}/preview',json={'revision':0,'mapping':MAPPING,'source_region':REGION,'reuse_layout':True}).status_code==200
    assert client.post(f'/api/usage/{identifier}/approve',json={'revision':1,'acknowledge':True}).status_code==200
    original=client.get(f'/api/sources/{inspected.json()["document_id"]}')
    assert original.content==raw and 'attachment' in original.headers['content-disposition']
    assert client.get('/api/overview').json()['total_cents']==0


def test_archive_duplicate_members_and_dimension_lies_are_not_trusted(usage):
    raw=xlsx()
    output=BytesIO()
    with ZipFile(BytesIO(raw)) as old,ZipFile(output,'w',ZIP_DEFLATED) as new:
        for name in old.namelist():new.writestr(name,old.read(name))
        with pytest.warns(UserWarning,match='Duplicate name'):
            new.writestr('xl/workbook.xml',old.read('xl/workbook.xml'))
    with pytest.raises(ValueError,match='DUPLICATE_MEMBER'):
        usage.import_file('duplicated.xlsx',output.getvalue())
    understated=rewrite(raw,'xl/worksheets/sheet1.xml',lambda v:v.replace(b'<dimension ref="B1:G6"/>',b'<dimension ref="A1:A1"/>'))
    identifier=ready_xlsx(usage,understated)
    assert usage.detail(identifier)['preview']['reading_count']==3


def test_native_1904_epoch_is_explicit_and_early_1900_dates_fail(usage):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    def native(b):
        b.epoch=CALENDAR_MAC_1904
        for r in range(4,7):
            for c in ('C','D'):
                cell=b['Water export'][f'{c}{r}']
                cell.value=datetime.fromisoformat(cell.value).replace(tzinfo=None)
                cell.number_format='yyyy-mm-dd hh:mm:ss'
    identifier=ready_xlsx(usage,xlsx(alter=native),mapping={**MAPPING,'timezone':'America/New_York'})
    data=usage.detail(identifier)['preview']
    assert data['provenance']['date_system']=='1904'
    assert data['readings'][0]['start_utc']==1785556800
    def early(b):
        cell=b['Water export']['C4'];cell.value=60;cell.number_format='yyyy-mm-dd hh:mm:ss'
    with pytest.raises(ValueError,match='DATE_OR_TIME_AMBIGUOUS'):
        usage.import_file('invalid-1900-calendar.xlsx',xlsx(alter=early))


def test_fixed_subregion_reuse_requires_unchanged_sheet_shape(usage):
    def footer(b):b['Water export']['B8']='Fictional footer outside selected region'
    first=ready_xlsx(usage,xlsx(alter=footer))
    approve(usage,first)
    raw=RAW.replace(b'2026-08-01',b'2026-09-01')
    second=usage.import_file('same-region.xlsx',xlsx(raw,footer))['import_id']
    assert usage.detail(second)['source_region']==REGION
    def shifted_footer(b):b['Water export']['B9']='Fictional moved footer'
    changed=usage.import_file('changed-region.xlsx',xlsx(raw,shifted_footer))['import_id']
    assert usage.detail(changed)['suggested_mapping'] is None


def test_provenance_preview_change_disables_stale_approval_and_old_bytes_remain(usage):
    identifier=ready_xlsx(usage)
    usage.preview(identifier,{'revision':1,'mapping':MAPPING,'source_region':{**REGION,'end_row':5},'reuse_layout':False})
    with pytest.raises(ValueError,match='CHANGED_REOPEN_PREVIEW'):
        approve(usage,identifier)
    assert usage.listing()['active_reading_count']==0
    data=usage.detail(identifier)
    assert data['revision']==2 and data['preview']['reading_count']==2
    with usage.store.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM usage_previews WHERE import_id=?',(identifier,)).fetchone()[0]==2
    assert approve(usage,identifier,revision=2)['reading_count']==2
    assert check(usage.store)['audit']=='ok'


def test_mapping_uses_pinned_timezone_data_even_with_empty_or_poisoned_system_path(usage,tmp_path):
    import zoneinfo
    from importlib.resources import files
    from utilityos.usage import source_timezone,timestamp
    original_path=zoneinfo.TZPATH
    poisoned=tmp_path/'poisoned-system-zoneinfo'
    (poisoned/'America').mkdir(parents=True)
    # A conflicting OS database labels UTC as New York. Application mapping
    # must still use the reviewed package's actual Eastern offset.
    (poisoned/'America/New_York').write_bytes(files('tzdata.zoneinfo').joinpath('Etc','UTC').read_bytes())
    try:
        for path in ((),(str(poisoned),)):
            zoneinfo.reset_tzpath(path)
            zoneinfo.ZoneInfo.clear_cache()
            zone=source_timezone('America/New_York')
            assert timestamp('2026-08-01T00:00:00',zone)==timestamp('2026-08-01T04:00:00Z',None)
            for value in ('2026-11-01T01:30:00','2026-03-08T02:30:00'):
                with pytest.raises(ValueError,match='DST_AMBIGUOUS_OR_NONEXISTENT'):
                    timestamp(value,zone)
            assert timestamp('2026-11-01T01:30:00-05:00',zone)==timestamp('2026-11-01T06:30:00Z',None)
        def native(b):
            for r in range(4,7):
                for c in ('C','D'):
                    cell=b['Water export'][f'{c}{r}'];cell.value=datetime.fromisoformat(cell.value).replace(tzinfo=None)
                    cell.number_format='yyyy-mm-dd hh:mm:ss'
        identifier=ready_xlsx(usage,xlsx(alter=native),mapping={**MAPPING,'timezone':'America/New_York'})
        data=usage.detail(identifier)['preview']
        assert data['readings'][0]['start_utc']==1785556800
        assert data['provenance']['timezone_database']=={'kind':'packaged_tzdata','version':'2026.4','zone':'America/New_York'}
        approve(usage,identifier)
    finally:
        zoneinfo.reset_tzpath(original_path)
        zoneinfo.ZoneInfo.clear_cache()


@pytest.mark.parametrize('key',['../UTC','America/../UTC','/etc/passwd','America\\New_York','America//New_York','__init__.py'])
def test_timezone_keys_cannot_escape_packaged_database(key):
    from utilityos.usage import source_timezone
    with pytest.raises(ValueError,match='USAGE_IANA_TIMEZONE_REQUIRED'):
        source_timezone(key)


def test_mismatched_or_missing_packaged_timezone_database_fails_closed(monkeypatch):
    from utilityos.usage import source_timezone
    import tzdata
    monkeypatch.setattr(tzdata,'__version__','UNREVIEWED_VERSION')
    with pytest.raises(ValueError,match='REVIEWED_TIMEZONE_DATABASE_REQUIRED'):
        source_timezone('America/New_York')
    monkeypatch.delattr(tzdata,'__version__')
    with pytest.raises(ValueError,match='REVIEWED_TIMEZONE_DATABASE_REQUIRED'):
        source_timezone('America/New_York')
