"""Bounded values-only XLSX inspection. Never save, recalculate or refresh a workbook."""
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import posixpath
import re
from zlib import error as ZipCompressionError
from zipfile import ZipFile, BadZipFile, ZIP_STORED, ZIP_DEFLATED
from defusedxml.ElementTree import iterparse
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from .parsers import ValidationError

MAX_BYTES = 8 * 1024 * 1024
MAX_EXPANDED = 32 * 1024 * 1024
MAX_CELLS = 150_030
MAX_SHEET_ROWS = 6000
MAX_SHEET_COLUMNS = 50
MAIN = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
REGION_FIELDS = {'sheet', 'header_row', 'first_column', 'last_column', 'end_row'}


def _check_text(value):
    if len(value) > 250 or any(ord(c) < 32 for c in value):
        raise ValidationError('USAGE_CELL_TOO_LONG_OR_CONTROL_CHARACTER')
    if value.lstrip().startswith(('=', '+', '-', '@')):
        raise ValidationError('USAGE_FORMULA_LIKE_CELL_REJECTED')
    return value


def _preflight(raw):
    if not raw or len(raw) > MAX_BYTES:
        raise ValidationError('XLSX_REQUIRES_1_TO_8_MB')
    if raw.startswith(b'\xd0\xcf\x11\xe0'):
        raise ValidationError('XLSX_ENCRYPTED_OR_LEGACY_XLS_UNSUPPORTED')
    trees = {}
    node_count = 0
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 200 or sum(e.file_size for e in entries) > MAX_EXPANDED:
            raise ValidationError('XLSX_ARCHIVE_RESOURCE_LIMIT')
        names = set()
        for entry in entries:
            name = entry.filename
            if (name.casefold() in names or name.startswith('/') or '\\' in name or
                    any(part in {'', '.', '..'} for part in name.split('/'))):
                raise ValidationError('XLSX_ARCHIVE_UNSAFE_OR_DUPLICATE_MEMBER')
            names.add(name.casefold())
            if (entry.flag_bits & 1 or entry.compress_type not in {ZIP_STORED, ZIP_DEFLATED}):
                raise ValidationError('XLSX_ENCRYPTION_OR_COMPRESSION_UNSUPPORTED')
            if entry.file_size > MAX_BYTES or entry.file_size > max(1024, entry.compress_size * 200):
                raise ValidationError('XLSX_ARCHIVE_RESOURCE_LIMIT')
            lower = name.lower()
            if (not lower.endswith(('.xml', '.rels')) or any(token in lower for token in
                    ('vba', 'macrosheet', 'externallink', 'connections', 'querytable', 'pivot', 'embedding', 'activex', 'calcchain'))):
                raise ValidationError('XLSX_ACTIVE_OR_UNSUPPORTED_CONTENT')
            parser = iterparse(BytesIO(archive.read(entry)), events=('start','end'),
                               forbid_dtd=True, forbid_entities=True, forbid_external=True)
            depth = 0
            for action, node in parser:
                depth += 1 if action == 'start' else -1
                if action == 'start':
                    node_count += 1
                if node_count > 650_000 or depth > 64:
                    raise ValidationError('XLSX_XML_NODE_RESOURCE_LIMIT')
            tree = parser.root
            trees[name] = tree
            for node in tree.iter():
                local = node.tag.rsplit('}', 1)[-1]
                if local in {'f', 'formula', 'formula1', 'formula2'}:
                    raise ValidationError('XLSX_FORMULA_OR_CACHED_FORMULA_REJECTED')
                if local == 'Relationship' and (node.get('TargetMode') == 'External' or
                        re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', node.get('Target', ''))):
                    raise ValidationError('XLSX_EXTERNAL_RELATIONSHIP_REJECTED')
                if local in {'Override', 'Default'} and any(token in node.get('ContentType', '').lower()
                        for token in ('macro', 'vba', 'activex', 'oleobject')):
                    raise ValidationError('XLSX_ACTIVE_OR_UNSUPPORTED_CONTENT')
        required = {'[Content_Types].xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels'}
        if not required.issubset(trees):
            raise ValidationError('XLSX_WORKBOOK_PARTS_REQUIRED')
        rels = {r.get('Id'): r.get('Target', '') for r in trees['xl/_rels/workbook.xml.rels']}
        sheets = []
        total_cells = 0
        for sheet in trees['xl/workbook.xml'].findall(MAIN+'sheets/'+MAIN+'sheet'):
            if sheet.get('state', 'visible') != 'visible':
                raise ValidationError('XLSX_HIDDEN_SHEETS_REQUIRE_VISIBLE_VALUES_EXPORT')
            name = sheet.get('name', '')
            if not name or len(name) > 100:
                raise ValidationError('XLSX_SHEET_NAME_INVALID')
            target = rels.get(sheet.get(REL+'id'), '')
            path = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/'+target)
            if not path.startswith('xl/worksheets/') or path not in trees:
                raise ValidationError('XLSX_WORKSHEET_RELATIONSHIP_UNSUPPORTED')
            tree = trees[path]
            cells = {}
            max_row = max_col = 0
            previous_row = 0
            for row in tree.findall(MAIN+'sheetData/'+MAIN+'row'):
                row_number = int(row.get('r', '0'))
                if not previous_row < row_number <= MAX_SHEET_ROWS:
                    raise ValidationError('XLSX_ROW_ORDER_OR_RESOURCE_LIMIT')
                previous_row = row_number
                previous_col = 0
                for cell in row.findall(MAIN+'c'):
                    address = cell.get('r', '')
                    r, c = coordinate_to_tuple(address)
                    if (r != row_number or not previous_col < c <= MAX_SHEET_COLUMNS or address in cells):
                        raise ValidationError('XLSX_CELL_ORDER_OR_RESOURCE_LIMIT')
                    previous_col = c
                    value = cell.find(MAIN+'v')
                    stored = value.text if value is not None and value.text is not None else ''
                    if len(stored) > 250:
                        raise ValidationError('USAGE_CELL_TOO_LONG_OR_CONTROL_CHARACTER')
                    cells[address] = {'raw_xml_value': stored, 'cell_type': cell.get('t', 'n')}
                    max_row, max_col = max(max_row, r), max(max_col, c)
                    total_cells += 1
                    if total_cells > MAX_CELLS:
                        raise ValidationError('XLSX_CELL_RESOURCE_LIMIT')
            merged = [range_boundaries(n.get('ref', '')) for n in tree.findall(MAIN+'mergeCells/'+MAIN+'mergeCell')]
            hidden_rows = [int(n.get('r')) for n in tree.findall(MAIN+'sheetData/'+MAIN+'row') if n.get('hidden') in {'1', 'true'}]
            hidden_cols = [(int(n.get('min')), int(n.get('max'))) for n in tree.findall(MAIN+'cols/'+MAIN+'col') if n.get('hidden') in {'1', 'true'}]
            sheets.append({'name': name, 'max_row': max_row, 'max_column': max_col, 'cells': cells,
                           'merged': merged, 'hidden_rows': hidden_rows, 'hidden_columns': hidden_cols})
        if not 1 <= len(sheets) <= 10 or len({s['name'] for s in sheets}) != len(sheets):
            raise ValidationError('XLSX_REQUIRES_1_TO_10_UNIQUE_VISIBLE_SHEETS')
        return sheets


def workbook(raw):
    """Return bounded plain values with exact XML numeric lexemes and cell locations."""
    book = None
    try:
        sheets = _preflight(raw)
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False, keep_vba=False)
        for sheet in sheets:
            values = {}
            ws = book[sheet['name']]
            for row in ws.iter_rows(min_row=1, max_row=max(1, sheet['max_row']), min_col=1, max_col=max(1, sheet['max_column'])):
                for cell in row:
                    if cell.value is None:
                        continue
                    address = cell.coordinate
                    meta = sheet['cells'][address]
                    value = cell.value
                    if cell.data_type in {'f', 'e', 'b'}:
                        raise ValidationError('XLSX_FORMULA_ERROR_OR_BOOLEAN_UNSUPPORTED')
                    if isinstance(value, datetime):
                        # Excel datetimes have no zone. Require an explicit time format;
                        # date-only or duration formats cannot establish interval bounds.
                        if (value.microsecond or not re.search(r'[hs]', cell.number_format.lower()) or
                                (book.epoch.year == 1899 and meta['cell_type'] == 'n' and Decimal(meta['raw_xml_value']) < 61)):
                            raise ValidationError('XLSX_DATE_OR_TIME_AMBIGUOUS_USE_ISO_DATETIME')
                        text = value.isoformat(timespec='seconds')
                    elif isinstance(value, str):
                        text = value
                    elif isinstance(value, (int, float)):
                        # Do not round a provider's decimal through a binary float.
                        text = meta['raw_xml_value']
                    else:
                        raise ValidationError('XLSX_DATE_OR_TIME_AMBIGUOUS_USE_ISO_DATETIME')
                    values[address] = _check_text(text)
                    meta['raw_value'] = text if meta['cell_type'] not in {'n', 'd'} else meta['raw_xml_value']
                    meta['number_format'] = cell.number_format[:100]
            sheet['values'] = values
        return {'sheets': sheets, 'date_system': '1904' if book.epoch.year == 1904 else '1900'}
    except ValidationError:
        raise
    except (BadZipFile, KeyError, ValueError, TypeError, AttributeError, IndexError, OverflowError,
            OSError, RuntimeError, ParseError, DefusedXmlException, InvalidOperation, ZipCompressionError):
        raise ValidationError('XLSX_INVALID_OR_UNSUPPORTED_WORKBOOK') from None
    finally:
        if book is not None:
            book.close()


def inspect_workbook(book):
    """A bounded sheet inventory plus candidate header rows, never auto-selected."""
    result = []
    for sheet in book['sheets']:
        candidates = []
        for number in range(1, min(sheet['max_row'], 25)+1):
            populated = [(column, sheet['values'].get(f'{get_column_letter(column)}{number}', ''))
                         for column in range(1, sheet['max_column']+1)]
            populated = [(c, value) for c, value in populated if value]
            if 2 <= len(populated) <= 30 and len({v for _, v in populated}) == len(populated):
                candidates.append({'header_row': number, 'first_column': populated[0][0],
                                   'last_column': populated[-1][0], 'values': [v for _, v in populated]})
        result.append({'name': sheet['name'], 'max_row': sheet['max_row'], 'max_column': sheet['max_column'],
                       'candidate_headers': candidates[:5]})
    return result


def table(book, region):
    if not isinstance(region, dict) or set(region) != REGION_FIELDS:
        raise ValidationError('XLSX_EXPLICIT_SHEET_AND_REGION_REQUIRED')
    if not isinstance(region['sheet'], str) or any(type(region[k]) is not int for k in REGION_FIELDS-{'sheet'}):
        raise ValidationError('XLSX_REGION_INVALID')
    sheet = next((s for s in book['sheets'] if s['name'] == region['sheet']), None)
    header, end, first, last = (region[k] for k in ('header_row', 'end_row', 'first_column', 'last_column'))
    if (not sheet or not 1 <= header < end <= sheet['max_row'] or end-header > 5000 or
            not 1 <= first < last <= sheet['max_column'] or last-first+1 > 30):
        raise ValidationError('XLSX_REGION_REQUIRES_HEADER_AND_1_TO_5000_DATA_ROWS')
    if (any(header <= r <= end for r in sheet['hidden_rows']) or
            any(a <= last and b >= first for a, b in sheet['hidden_columns']) or
            any(a <= last and c >= first and b <= end and d >= header for a, b, c, d in sheet['merged'])):
        raise ValidationError('XLSX_REGION_CONTAINS_MERGED_OR_HIDDEN_CELLS')
    headers = [sheet['values'].get(f'{get_column_letter(c)}{header}', '') for c in range(first, last+1)]
    if len(set(headers)) != len(headers) or any(not h.strip() or h != h.strip() for h in headers):
        raise ValidationError('USAGE_HEADERS_MUST_BE_UNIQUE_NONEMPTY')
    rows, cells = [], []
    for r in range(header+1, end+1):
        row, evidence = {}, {}
        for c, heading in zip(range(first, last+1), headers):
            address = f'{get_column_letter(c)}{r}'
            row[heading] = sheet['values'].get(address, '')
            evidence[heading] = {**sheet['cells'].get(address, {'raw_value': '', 'raw_xml_value': '', 'cell_type': 'empty'}),
                                 'cell': address, 'row': r, 'column': c}
        if not any(row.values()):
            raise ValidationError('XLSX_BLANK_DATA_ROW_REQUIRES_REGION_REVIEW')
        rows.append(row)
        cells.append(evidence)
    return {'headers': headers, 'rows': rows, 'cells': cells, 'region': region,
            'date_system': book['date_system'], 'format': 'xlsx'}
